import pptx
from pptx.dml.color import RGBColor
from pptx.util import Pt

def set_run_style(run, font_name="Calibri", size_pt=None, bold=None, italic=None, color_rgb=None):
    if font_name:
        run.font.name = font_name
    if size_pt:
        run.font.size = Pt(size_pt)
    if bold is not None:
        run.font.bold = bold
    if italic is not None:
        run.font.italic = italic
    if color_rgb:
        run.font.color.rgb = RGBColor(*color_rgb)

def replace_single_para(tf_shape, new_text, font_name="Calibri", size_pt=16.5, bold=None, italic=None, color_rgb=(30, 38, 56)):
    """Replace a text frame that has a single paragraph."""
    tf = tf_shape.text_frame
    tf.clear()
    p = tf.paragraphs[0]
    r = p.add_run()
    r.text = new_text
    set_run_style(r, font_name=font_name, size_pt=size_pt, bold=bold, italic=italic, color_rgb=color_rgb)

def set_two_para_body(tf_shape, p0_text, p1_text,
                      p0_kw=None, p1_kw=None):
    """Set a text frame body with exactly two paragraphs."""
    if p0_kw is None:
        p0_kw = dict(font_name="Calibri", size_pt=16.5, italic=True, color_rgb=(86, 96, 116))
    if p1_kw is None:
        p1_kw = dict(font_name="Calibri", size_pt=16.5, color_rgb=(30, 38, 56))
    tf = tf_shape.text_frame
    tf.clear()
    # P0
    p0 = tf.paragraphs[0]
    r0 = p0.add_run()
    r0.text = p0_text
    set_run_style(r0, **p0_kw)
    # P1
    p1 = tf.add_paragraph()
    r1 = p1.add_run()
    r1.text = p1_text
    set_run_style(r1, **p1_kw)

def apply_all_fixes():
    filepath = "Copy of SIH26177_v6.pptx.pptx"
    prs = pptx.Presentation(filepath)

    # =====================================================================
    # FIX 2: Slide 2 — Feature card 5: Remove incorrect Bouabdallah attribution
    #         for VIO. TextBox 162 is shape index to find by text content.
    # =====================================================================
    slide2 = prs.slides[1]
    for shape in slide2.shapes:
        if shape.has_text_frame and "Bouabdallah 2007" in shape.text_frame.text:
            tf = shape.text_frame
            tf.clear()
            p = tf.paragraphs[0]
            r = p.add_run()
            r.text = "Visual-Inertial Odometry overrides magnetometer in collapsed steel rubble (IEEE)."
            set_run_style(r, font_name="Calibri", size_pt=16.5, color_rgb=(30, 38, 56))
            print(f"  FIX 2 applied: {shape.name}")
            break

    # =====================================================================
    # FIX 3: Slide 3 — Step 6 workflow phrasing fix
    # =====================================================================
    slide3 = prs.slides[2]
    for shape in slide3.shapes:
        if shape.has_text_frame and "can use the Dijkstra" in shape.text_frame.text:
            tf = shape.text_frame
            tf.clear()
            p = tf.paragraphs[0]
            r = p.add_run()
            r.text = ("Offline GCS computes priority triage queue and charts a hazard-aware Dijkstra "
                      "safe walking corridor avoiding live wires directly to casualty.")
            set_run_style(r, font_name="Calibri", size_pt=16.5, color_rgb=(30, 38, 56))
            print(f"  FIX 3 applied: {shape.name}")
            break

    # =====================================================================
    # FIX 4: Slide 4 — Financial feasibility: Fix prototype cost + align with Slide 5
    # =====================================================================
    slide4 = prs.slides[3]
    for shape in slide4.shapes:
        if shape.has_text_frame and "20,000" in shape.text_frame.text:
            tf = shape.text_frame
            tf.clear()
            p0 = tf.paragraphs[0]
            r0 = p0.add_run()
            r0.text = "Prototype Build: ₹35,800 (Calibrated Functional Prototype)."
            set_run_style(r0, font_name="Calibri", size_pt=16.5, color_rgb=(30, 38, 56))
            p1 = tf.add_paragraph()
            r1 = p1.add_run()
            r1.text = "Production BOM: ₹85,000 – ₹1,20,000 per unit (10x cheaper than commercial enterprise drones)."
            set_run_style(r1, font_name="Calibri", size_pt=16.5, color_rgb=(30, 38, 56))
            print(f"  FIX 4 applied: {shape.name}")
            break

    # =====================================================================
    # FIX 5: Slide 6 — Presentation video blank box: Add placeholder text
    # =====================================================================
    slide6 = prs.slides[5]
    # Find the autoshape placeholder box near "Presentation Video:"
    # It is AutoShape 71 (shape index 29) - an empty autoshape box
    for shape in slide6.shapes:
        if shape.has_text_frame and shape.text_frame.text.strip() == "":
            # Check if this is in the general video area (position-based check)
            # The video box is at approx left=10767060, top=4200000-5500000 range
            if hasattr(shape, 'left') and shape.left > 9000000 and shape.top < 6000000:
                tf = shape.text_frame
                tf.clear()
                p = tf.paragraphs[0]
                r = p.add_run()
                r.text = "Video link: SIH Grand Finale Submission Portal"
                set_run_style(r, font_name="Calibri", size_pt=13, italic=True, color_rgb=(120, 130, 150))
                print(f"  FIX 5 applied: {shape.name} at left={shape.left}, top={shape.top}")
                break

    prs.save(filepath)
    print(f"\nAll fixes saved successfully to {filepath}")

if __name__ == "__main__":
    print("Applying targeted fixes...")
    apply_all_fixes()
