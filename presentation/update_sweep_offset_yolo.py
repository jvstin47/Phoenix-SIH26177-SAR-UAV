"""
Update Slide 2 and Slide 3 to:
1. Emphasize Autonomous Sweep mechanics (Boustrophedon 60% sidelap / 70% overlap / ToF terrain adapt).
2. Emphasize Ground Offset Geolocation (2D bbox offset -> 3D Tait-Bryan DCM matrix -> GPS pin on the fly).
3. Explicitly state YOLO's dual role:
   - Primary: Survivor identification (37°C body heat contour extraction).
   - Secondary: Dynamic obstacle & flight hazard avoidance (wires, structural collapse obstacles).
4. Reduce redundant slogans.
"""
import pptx
from pptx.dml.color import RGBColor
from pptx.util import Pt

FILEPATH = "Copy of SIH26177_v6.pptx.pptx"

def set_para_text(p, text, font_name="Calibri", size_pt=15.25, color_rgb=(30, 38, 56), bold=False, italic=False):
    p.text = text
    if p.runs:
        r = p.runs[0]
        r.font.name = font_name
        r.font.size = Pt(size_pt)
        r.font.color.rgb = RGBColor(*color_rgb)
        r.font.bold = bold
        r.font.italic = italic

def update_presentation():
    prs = pptx.Presentation(FILEPATH)

    # =========================================================================
    # SLIDE 2: Card 2 (TextBox 159) - YOLO survivor + obstacle avoidance
    # =========================================================================
    slide2 = prs.slides[1]
    for s in slide2.shapes:
        if s.name == "TextBox 159":
            tf = s.text_frame
            tf.clear()
            p = tf.paragraphs[0]
            r = p.add_run()
            r.text = "INT8 YOLO identifies victims & flight obstacles on Qualcomm NPU; preserves 90% power for rotors."
            r.font.name = "Calibri"
            r.font.size = Pt(16.5)
            r.font.color.rgb = RGBColor(30, 38, 56)
            print("Slide 2 Card 2 updated (TextBox 159)")
            break

    # =========================================================================
    # SLIDE 3: Technical Approach
    # =========================================================================
    slide3 = prs.slides[2]

    # Step 2: Autonomous Sweep mechanics (TextBox 46)
    for s in slide3.shapes:
        if s.name == "TextBox 46":
            tf = s.text_frame
            tf.clear()
            p = tf.paragraphs[0]
            r = p.add_run()
            r.text = "Boustrophedon decomposition with 60% sidelap & 70% overlap; ToF altimeter maintains constant GSD over sloped terrain."
            r.font.name = "Calibri"
            r.font.size = Pt(15.25)
            r.font.color.rgb = RGBColor(30, 38, 56)
            print("Slide 3 Step 2 updated (TextBox 46 - Sweep mechanics)")
            break

    # Step 3: YOLO Survivor ID + Obstacle Avoidance (TextBox 56)
    for s in slide3.shapes:
        if s.name == "TextBox 56":
            tf = s.text_frame
            tf.clear()
            p = tf.paragraphs[0]
            r = p.add_run()
            r.text = "Dual 4K RGB + LWIR Thermal fusion; edge YOLO identifies survivor heat contours (37°C) & flags obstacles for flight avoidance."
            r.font.name = "Calibri"
            r.font.size = Pt(15.25)
            r.font.color.rgb = RGBColor(30, 38, 56)
            print("Slide 3 Step 3 updated (TextBox 56 - YOLO survivor + obstacle avoidance)")
            break

    # Step 4: Offset Geolocation mechanics (TextBox 66)
    for s in slide3.shapes:
        if s.name == "TextBox 66":
            tf = s.text_frame
            tf.clear()
            p = tf.paragraphs[0]
            r = p.add_run()
            r.text = "Transforms 2D bbox offset via 3D Tait-Bryan DCM matrix & ToF altitude (ΔP = H_AGL · R·v) to pin GPS on the fly."
            r.font.name = "Calibri"
            r.font.size = Pt(15.25)
            r.font.color.rgb = RGBColor(30, 38, 56)
            print("Slide 3 Step 4 updated (TextBox 66 - Dynamic offset geolocation)")
            break

    # Mission Computer: Dual-Task YOLO (TextBox 112)
    for s in slide3.shapes:
        if s.name == "TextBox 112":
            tf = s.text_frame
            tf.clear()
            p0 = tf.paragraphs[0]
            r0 = p0.add_run()
            r0.text = "Dual-Task INT8 YOLO (Survivor Triage + Flight Obstacle Avoidance) // Ray-Caster"
            r0.font.name = "Calibri"
            r0.font.size = Pt(12.89)
            r0.font.color.rgb = RGBColor(228, 236, 247)

            p1 = tf.add_paragraph()
            r1 = p1.add_run()
            r1.text = "Semtech SX1262 LoRa Telemetry Driver (868 MHz)"
            r1.font.name = "Calibri"
            r1.font.size = Pt(12.89)
            r1.font.color.rgb = RGBColor(228, 236, 247)
            print("Slide 3 Mission Computer updated (TextBox 112)")
            break

    # Tech Stack: Edge AI & Vision (Group 130 > TextBox 132)
    for s in slide3.shapes:
        if s.shape_type == 6 and s.name == "Group 130":
            for sub in s.shapes:
                if sub.name == "TextBox 132":
                    tf = sub.text_frame
                    tf.clear()
                    p = tf.paragraphs[0]
                    r = p.add_run()
                    r.text = "YOLOv8n INT8 (Detection & Avoidance), SNPE, OpenCV"
                    r.font.name = "Calibri"
                    r.font.size = Pt(14)
                    r.font.color.rgb = RGBColor(30, 38, 56)
                    print("Slide 3 Tech Stack Edge AI updated (TextBox 132)")
                    break

    prs.save(FILEPATH)
    print(f"\nAll updates successfully saved to {FILEPATH}")

if __name__ == "__main__":
    update_presentation()
