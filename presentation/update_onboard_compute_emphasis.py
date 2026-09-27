"""
Update Slide 2, Slide 3, and Slide 4 to emphasize the powerful dual-brain onboard architecture:
- All heavy AI, frame registration, ray-casting, and triage happen ONBOARD the UAV.
- Ground operators do NOT need high-end GPUs, servers, or cellular connectivity.
- GCS acts as a lightweight, low-power terminal receiving 17-byte LoRa packets.
"""
import pptx
from pptx.dml.color import RGBColor
from pptx.util import Pt

FILEPATH = "Copy of SIH26177_v6.pptx.pptx"

def update_onboard_compute_emphasis():
    prs = pptx.Presentation(FILEPATH)

    # =========================================================================
    # SLIDE 2: Solution Pitch Box (TextBox 157)
    # =========================================================================
    slide2 = prs.slides[1]
    for s in slide2.shapes:
        if s.name == "TextBox 157":
            tf = s.text_frame
            # P1: 100% ONBOARD EDGE COMPUTE
            p1 = tf.paragraphs[1]
            p1.clear()
            r1_bold = p1.add_run()
            r1_bold.text = "100% ONBOARD EDGE COMPUTE: "
            r1_bold.font.name = "Calibri"
            r1_bold.font.size = Pt(16.25)
            r1_bold.font.bold = True
            r1_bold.font.color.rgb = RGBColor(30, 38, 56)

            r1_text = p1.add_run()
            r1_text.text = "All AI vision, sensor fusion & geolocation execute on UAV silicon; zero reliance on ground GPUs or cloud."
            r1_text.font.name = "Calibri"
            r1_text.font.size = Pt(16.25)
            r1_text.font.color.rgb = RGBColor(30, 38, 56)

            # P4: LIGHTWEIGHT OFFLINE GCS
            p4 = tf.paragraphs[4]
            p4.clear()
            r4_bold = p4.add_run()
            r4_bold.text = "LIGHTWEIGHT OFFLINE GCS: "
            r4_bold.font.name = "Calibri"
            r4_bold.font.size = Pt(16.25)
            r4_bold.font.bold = True
            r4_bold.font.color.rgb = RGBColor(30, 38, 56)

            r4_text = p4.add_run()
            r4_text.text = "Low-power field tablets receive 17-byte pins; charts priority triage queues & safe walking corridors."
            r4_text.font.name = "Calibri"
            r4_text.font.size = Pt(16.25)
            r4_text.font.color.rgb = RGBColor(30, 38, 56)

            print("Slide 2 TextBox 157 updated (Onboard Compute & Lightweight GCS)")
            break

    # =========================================================================
    # SLIDE 3: Steps 5 & 6 (TextBox 76 and TextBox 85)
    # =========================================================================
    slide3 = prs.slides[2]
    for s in slide3.shapes:
        if s.name == "TextBox 76":
            tf = s.text_frame
            tf.clear()
            p = tf.paragraphs[0]
            r = p.add_run()
            r.text = "Onboard Triage & LoRa Burst: UAV computes victim priority on-chip & packs coordinates into 17-byte binary frame; broadcasts over 5+ km."
            r.font.name = "Calibri"
            r.font.size = Pt(15.25)
            r.font.color.rgb = RGBColor(30, 38, 56)
            print("Slide 3 Step 5 updated (TextBox 76 - Onboard Triage)")
            break

    for s in slide3.shapes:
        if s.name == "TextBox 85":
            tf = s.text_frame
            tf.clear()
            p = tf.paragraphs[0]
            r = p.add_run()
            r.text = "Lightweight Offline GCS: Low-power tablet receives 17-byte pins; charts hazard-aware Dijkstra safe walking corridor avoiding live wires."
            r.font.name = "Calibri"
            r.font.size = Pt(15.25)
            r.font.color.rgb = RGBColor(30, 38, 56)
            print("Slide 3 Step 6 updated (TextBox 85 - Lightweight GCS)")
            break

    # =========================================================================
    # SLIDE 4: Operational Feasibility (TextBox 46)
    # =========================================================================
    slide4 = prs.slides[3]
    for s in slide4.shapes:
        if s.name == "TextBox 46":
            tf = s.text_frame
            tf.clear()

            p0 = tf.paragraphs[0]
            r0 = p0.add_run()
            r0.text = "Setup Time: < 3 minutes from rugged backpack unboxing to flight."
            r0.font.name = "Calibri"
            r0.font.size = Pt(15.5)
            r0.font.color.rgb = RGBColor(30, 38, 56)

            p1 = tf.add_paragraph()
            r1 = p1.add_run()
            r1.text = "Quick-Swap Battery: Dual 4S LiPo tray for continuous sortie coverage."
            r1.font.name = "Calibri"
            r1.font.size = Pt(15.5)
            r1.font.color.rgb = RGBColor(30, 38, 56)

            p2 = tf.add_paragraph()
            r2 = p2.add_run()
            r2.text = "Zero Ground Compute Burden: 100% onboard processing enables deployment via low-power tablets or handheld LoRa pagers."
            r2.font.name = "Calibri"
            r2.font.size = Pt(15.5)
            r2.font.color.rgb = RGBColor(30, 38, 56)

            print("Slide 4 Operational Feasibility updated (TextBox 46 - Zero Ground Compute Burden)")
            break

    prs.save(FILEPATH)
    print(f"\nAll onboard compute emphasis updates saved to {FILEPATH}")

if __name__ == "__main__":
    update_onboard_compute_emphasis()
